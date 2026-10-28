import csv
import hashlib
import io
import json
import zipfile

from fastapi import APIRouter, Depends, File, Query, UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from .common import csv_response, mutate, rows, validate_lead
from .config import settings
from .db import get_db
from .leads import create_lead
from .models import ImportRun, Lead
from .schemas import LeadIn
from .security import current, fail, phone

router = APIRouter(prefix="/api/imports", tags=["Imports"])
ALIASES = {
    "mobile": "phone",
    "contact number": "phone",
    "mobile number": "phone",
    "phone number": "phone",
    "customer": "name",
    "customer name": "name",
    "full name": "name",
    "customer/full name": "name",
}
ALLOWED = {"name", "phone", "city", "state", "source", "category", "interests", "notes", "stage", "custom", "details"}


async def bounded_upload(file, max_bytes=None):
    limit = settings.max_upload if max_bytes is None else max_bytes
    data = await file.read(limit + 1)
    if len(data) > limit:
        fail("file_too_large", f"Maximum file size is {limit // (1024 * 1024)} MiB", 413)
    if not data:
        fail("empty_file", "The file is empty", 422)
    return data


def parse(data, filename):
    try:
        if filename.lower().endswith(".csv"):
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary data is not CSV")
            reader = csv.reader(io.StringIO(text), strict=True)
            matrix = []
            for row in reader:
                matrix.append(row)
                if len(matrix) > 10001:
                    fail("too_many_rows", "Maximum 10,000 data rows per import", 422)
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    fail("expanded_file_too_large", "The expanded workbook exceeds the safe size limit", 413)
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            sheet = workbook.active
            if sheet.max_row and sheet.max_row > 10001:
                fail("too_many_rows", "Maximum 10,000 data rows per import", 422)
            if sheet.max_column and sheet.max_column > 100:
                fail("too_many_columns", "Maximum 100 columns", 422)
            matrix = []
            for row in sheet.iter_rows(values_only=True):
                matrix.append(list(row))
                if len(matrix) > 10001:
                    fail("too_many_rows", "Maximum 10,000 data rows per import", 422)
            workbook.close()
        else:
            fail("unsupported_file", "Save this file as .xlsx or UTF-8 .csv; legacy .xls is unsupported", 422)
    except Exception as exc:
        from fastapi import HTTPException

        if isinstance(exc, HTTPException):
            raise
        fail("malformed_file", "Could not read this file. Upload a valid UTF-8 CSV or XLSX workbook", 422)
    if not matrix:
        fail("empty_file", "The file is empty", 422)
    headers = [ALIASES.get(str(h or "").strip().lower(), str(h or "").strip().lower()) for h in matrix[0]]
    if not headers or len(headers) > 100 or "name" not in headers or "phone" not in headers:
        fail("missing_columns", "The header must include name and phone", 422)
    if len(headers) != len(set(headers)) or "" in headers:
        fail("duplicate_headers", "Headers must be unique, named columns, including aliases", 422)
    forbidden = set(headers) & {"owner", "owner_id", "assignee", "assigned_to", "org_id", "user_id"}
    if forbidden:
        fail(
            "forged_ownership", "Remove owner, assignee and organization columns. Ownership is set by your account", 422
        )
    if set(headers) - ALLOWED:
        fail("unknown_columns", "Unrecognized columns: " + ", ".join(sorted(set(headers) - ALLOWED)), 422)
    return headers, matrix[1:]


@router.get("/template")
def template(ctx=Depends(current)):
    return csv_response(
        [], ["name", "phone", "city", "state", "source", "category", "interests", "notes"], "lead-template.csv"
    )


@router.post("")
async def import_leads(
    file: UploadFile = File(),
    commit: bool = Query(False),
    ctx=Depends(current),
    db: Session = Depends(get_db, scope="function"),
):
    mutate(db, ctx)
    raw = await bounded_upload(file)
    digest = hashlib.sha256(raw).hexdigest()
    existing = db.scalar(rows(db, ImportRun, ctx.org_id).where(ImportRun.user_id == ctx.id, ImportRun.digest == digest))
    if existing:
        return {**existing.result, "replayed": True, "import_id": existing.id}
    headers, matrix = parse(raw, file.filename or "")
    existing_phones = set(db.scalars(select(Lead.phone).where(Lead.org_id == ctx.org_id, Lead.deleted.is_(False))))
    seen, accepted, errors, duplicates, blanks = set(), [], [], 0, 0
    for number, cells in enumerate(matrix, 2):
        if not any(v is not None and str(v).strip() for v in cells):
            blanks += 1
            continue
        if len(cells) != len(headers):
            # Excel can omit trailing empty cells; CSV needs an exact shape.
            if (file.filename or "").lower().endswith(".xlsx") and len(cells) < len(headers):
                cells += [None] * (len(headers) - len(cells))
            else:
                errors.append({"row": number, "column": "row", "message": "Column count does not match header"})
                continue
        item = dict(zip(headers, cells))
        try:
            for key, value in list(item.items()):
                if value is None or value == "":
                    item.pop(key)
                elif key in ["custom", "details"]:
                    item[key] = json.loads(str(value))
                elif key == "interests":
                    item[key] = [v.strip() for v in str(value).split("|") if v.strip()]
                elif key == "phone":
                    item[key] = phone(value)
                else:
                    item[key] = str(value).strip()
            payload = LeadIn(**item)
            validate_lead(db, ctx, payload.model_dump())
            if payload.phone in seen or payload.phone in existing_phones:
                duplicates += 1
                errors.append(
                    {
                        "row": number,
                        "column": "phone",
                        "message": "Duplicate phone in file or organization",
                        "duplicate": True,
                    }
                )
                continue
            seen.add(payload.phone)
            accepted.append(payload)
        except Exception as exc:
            if isinstance(exc, ValidationError):
                detail = exc.errors()[0]
                errors.append({"row": number, "column": str(detail["loc"][0]), "message": detail["msg"]})
            else:
                detail = getattr(exc, "detail", {})
                errors.append(
                    {
                        "row": number,
                        "column": "row",
                        "message": detail.get("message", "Invalid JSON or value")
                        if isinstance(detail, dict)
                        else "Invalid row",
                    }
                )
    result = {
        "accepted": len(accepted),
        "imported": 0,
        "duplicates": duplicates,
        "invalid": len(errors) - duplicates,
        "blank_rows": blanks,
        "skipped": len(errors) + blanks,
        "errors": errors,
        "preview": [p.model_dump(mode="json") for p in accepted[:10]],
        "owner": ctx.id if ctx.role in {"telecaller", "field_executive"} else None,
    }
    if commit:
        for data in accepted:
            create_lead(db, ctx, data)
        result["imported"] = len(accepted)
        run = ImportRun(org_id=ctx.org_id, user_id=ctx.id, digest=digest, result=result)
        db.add(run)
        db.flush()
        result["import_id"] = run.id
    return result


@router.get("/{import_id}/errors")
def import_errors(import_id: str, ctx=Depends(current), db: Session = Depends(get_db, scope="function")):
    from .common import get_scoped

    run = get_scoped(db, ImportRun, import_id, ctx)
    if ctx.role in {"telecaller", "field_executive"} and run.user_id != ctx.id:
        fail("not_found", "Import not found", 404)
    return csv_response(run.result.get("errors", []), ["row", "column", "message"], "import-errors.csv")
